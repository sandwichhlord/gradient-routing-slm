"""Build presentation/midterm_deck.pptx (10 slides, 16:9).  Run from anywhere: python presentation/build_deck.py

Every number comes from results/baseline, results/routed, docs/ or the configs; see the speaker notes.
"""
import os
import random

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.dml import MSO_LINE
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT = os.path.join(HERE, "midterm_deck.pptx")

# ---- palette: the three domain colours from our own plots drive the whole deck
NAVY, INK, MUTED, LIGHT, WHITE = "14213D", "1B1F2A", "5B6475", "F2F4F8", "FFFFFF"
STORY, MATH, CODE = "2A78D6", "EB6834", "1BAF7A"
AMBER, GREY, RED = "F2B134", "D5D9E0", "E34948"
HEAD, BODY = "Cambria", "Calibri"


def rgb(h):
    return RGBColor.from_string(h)


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def new_slide(dark=False):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(NAVY if dark else WHITE)
    return s


def text(slide, x, y, w, h, s, size=18, bold=False, color=INK, font=BODY, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, italic=False, margin=0.05, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    paras = s if isinstance(s, list) else [s]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        runs = p if isinstance(p, list) else [(p, {})]
        for t, o in runs:
            r = para.add_run()
            r.text = t
            r.font.name = o.get("font", font)
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", bold)
            r.font.italic = o.get("italic", italic)
            r.font.color.rgb = rgb(o.get("color", color))
    return tb


def box(slide, x, y, w, h, fill=LIGHT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08, lw=1.0, name=None):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        sh.name = name
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    return sh


def label_in(shape, s, size=18, bold=False, color=INK, font=BODY, align=PP_ALIGN.CENTER):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    paras = s if isinstance(s, list) else [s]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        r = para.add_run()
        r.text = p
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, font
        r.font.color.rgb = rgb(color)


def arrow(slide, x1, y1, x2, y2, color=MUTED, w=2.25, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    if dash:
        c.line.dash_style = MSO_LINE.DASH
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
    tail.set("type", "triangle")
    return c


def title(slide, s, n):
    text(slide, 0.6, 0.35, 12.1, 0.9, s, size=36, bold=True, color=NAVY, font=HEAD, anchor=MSO_ANCHOR.MIDDLE,
         name="Title")
    text(slide, 11.9, 7.0, 0.9, 0.3, f"{n} / 10", size=12, color=MUTED, align=PP_ALIGN.RIGHT)


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def picture(slide, rel, x, y, w=None, h=None):
    path = os.path.join(ROOT, rel)
    assert os.path.exists(path), path
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)


def chip(slide, x, y, w, h, s, fill, color=WHITE, size=14):
    c = box(slide, x, y, w, h, fill=fill, radius=0.5)
    label_in(c, s, size=size, bold=True, color=color)
    return c


# =============================================================================== 1. title
s = new_slide(dark=True)
for i in range(8):                       # the 8-block motif, 3 domain colours + grey
    col = {1: STORY, 4: MATH, 6: CODE}.get(i, "2B3A5E")
    box(s, 0.8 + i * 0.62, 0.9, 0.5, 0.5, fill=col, radius=0.18)
text(s, 0.8, 1.9, 11.7, 2.3, "Interpretable-by-Design Language Models via Activation-Based Gradient Routing",
     size=44, bold=True, color=WHITE, font=HEAD, name="Title")
text(s, 0.8, 4.35, 11.7, 0.6, "Minor Project  |  Mid-Term Evaluation", size=26, color="CADCFC")
text(s, 0.8, 5.0, 11.7, 0.5, "Department of CSE, PEC", size=22, color=WHITE)
text(s, 0.8, 5.65, 11.7, 0.5, "Team: [Name 1]  ·  [Name 2]  ·  [Name 3]  ·  [Name 4]", size=20, color=WHITE)
text(s, 0.8, 6.2, 11.7, 0.5, "Guide: [Guide name]    |    [Date]", size=20, color="CADCFC")
notes(s, """[0:00–1:00]  (Presenter 1)
Good morning. We are [names], and this is our minor project: "Interpretable-by-Design Language Models via Activation-Based Gradient Routing".
That is a long title, so here is the plain version. The AI language models behind chatbots are black boxes: we cannot see where a skill is stored inside them. Our project asks whether a model can be trained so that it organises itself into specialised parts, a bit like the regions of a brain, so that we can look inside it.
In the next ten minutes we will explain the problem, what others have done, our idea, how we test it, what we have already built and measured, and what is left to do.""")

# =============================================================================== 2. background
s = new_slide()
title(s, "A language model is a black box", 2)
text(s, 0.6, 1.35, 12.1, 1.0,
     [[("A language model ", {}), ("predicts the next word", {"bold": True}),
       (". It learns by adjusting millions of numbers called ", {}), ("weights", {"bold": True}), (".", {})]],
     size=22)
# left panel: brain with specialised regions
box(s, 0.6, 2.5, 5.9, 3.7, fill=LIGHT, name="Brain panel")
text(s, 0.8, 2.6, 5.5, 0.5, "A brain: specialised regions", size=22, bold=True, color=NAVY)
for (x, y, w, h, col, lab) in [(1.0, 3.3, 2.5, 1.35, STORY, "Language"), (3.65, 3.3, 2.5, 1.35, MATH, "Vision"),
                               (1.0, 4.8, 5.15, 1.15, CODE, "Movement")]:
    r = box(s, x, y, w, h, fill=col, radius=0.25)
    label_in(r, lab, size=20, bold=True, color=WHITE)
# right panel: model with mixed weights
box(s, 6.85, 2.5, 5.9, 3.7, fill=LIGHT, name="Model panel")
text(s, 7.05, 2.6, 5.5, 0.5, "A language model: no regions", size=22, bold=True, color=NAVY)
rnd = random.Random(7)
for _ in range(190):
    d = 0.17
    x, y = rnd.uniform(7.15, 12.2), rnd.uniform(3.3, 5.8)
    box(s, x, y, d, d, fill=rnd.choice([STORY, MATH, CODE]), shape=MSO_SHAPE.OVAL)
text(s, 0.6, 6.35, 12.1, 0.55, "Grammar, math and code are mixed across every weight.", size=20, color=INK,
     align=PP_ALIGN.CENTER)
notes(s, """[1:00–2:00]
A language model is a program that learns to predict the next word in a text. It does this by adjusting millions of numbers, called weights, while it reads lots of text. Our model learned from three kinds of text: children's stories, simple arithmetic, and Python code.
Here is the catch. In today's models, every skill is spread over all the weights, mixed together, like the coloured dots on the right.
Compare that with a human brain, on the left. A brain has specialised regions: one area for language, another for vision, another for movement. Doctors can tell which region is damaged from which ability is lost.
An AI model has no such regions. That is why we call it a black box: there is no place to look for "the math part".""")

# =============================================================================== 3. problem
s = new_slide()
title(s, "Why the black box is a problem", 3)
cards = [
    ("1", "Can't explain", "We cannot point to where an answer came from.", STORY),
    ("2", "Can't delete a skill", "Unlearning: removing one skill, such as unsafe knowledge, without starting over.", MATH),
    ("3", "Forgets old skills", "Catastrophic forgetting: learning something new overwrites old knowledge.", CODE),
]
for i, (n, head, body, col) in enumerate(cards):
    x = 0.5 + i * 4.2
    box(s, x, 1.5, 3.9, 3.7, fill=LIGHT, name=f"Card {n}")
    c = box(s, x + 0.3, 1.8, 0.8, 0.8, fill=col, shape=MSO_SHAPE.OVAL)
    label_in(c, n, size=28, bold=True, color=WHITE)
    text(s, x + 0.3, 2.8, 3.3, 0.5, head, size=24, bold=True, color=NAVY)
    text(s, x + 0.3, 3.4, 3.3, 1.7, body, size=18, color=INK)
q = box(s, 0.5, 5.6, 12.3, 1.2, fill=NAVY, name="Research question")
label_in(q, "Research question: can a model organise its own “brain regions” while it trains?", size=26,
         bold=True, color=WHITE)
notes(s, """[2:00–3:00]
This black-box design causes three concrete problems.
One: we cannot explain an answer. We cannot point to the part of the model that produced it.
Two: we cannot delete a skill. If a model learned something unsafe or private, "unlearning" it normally means retraining from scratch, because the knowledge is smeared everywhere.
Three: catastrophic forgetting. When we teach the model something new, it can overwrite what it knew before.
All three would be easier if skills lived in separate regions. So our research question is: can a model organise its own "brain regions" while it is training, without us telling it how?
▶ HANDOVER: Presenter 1 hands over to Presenter 2 for slide 4.""")

# =============================================================================== 4. comparison
s = new_slide()
title(s, "How our work differs from existing research", 4)
cols = [3.1, 3.5, 3.2, 2.5]
rows = [
    ("Approach", "What decides the regions?", "Goal", "Needs human labels?"),
    ("Mixture of Experts\n(Switch Transformer, Mixtral)", "A learned router, word by word", "Cheaper computing at large scale", "No"),
    ("Gradient Routing\n(Cloud et al., 2024)", "Humans assign data to regions", "Localise a skill, remove it later", "Yes"),
    ("EWC\n(Kirkpatrick et al., 2017)", "An importance score for each weight", "Don't forget old tasks", "Task boundaries"),
    ("Ours", "The model's own activity", "Self-organised regions we can switch off", "No"),
]
tbl = s.shapes.add_table(len(rows), 4, Inches(0.5), Inches(1.45), Inches(sum(cols)), Inches(3.95)).table
for j, w in enumerate(cols):
    tbl.columns[j].width = Inches(w)
for i, row in enumerate(rows):
    tbl.rows[i].height = Inches(0.55 if i == 0 else 0.85)
    for j, val in enumerate(row):
        cell = tbl.cell(i, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(NAVY if i == 0 else ("DCEBFA" if i == 4 else (WHITE if i % 2 else LIGHT)))
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = Inches(0.1)
        tf = cell.text_frame
        tf.word_wrap = True
        lines = val.split("\n")
        for k, ln in enumerate(lines):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            r = p.add_run()
            r.text = ln
            r.font.name = BODY
            r.font.size = Pt(16 if k == 0 or i == 0 else 14)
            r.font.bold = (i == 0) or (j == 0 and k == 0) or i == 4
            r.font.color.rgb = rgb(WHITE if i == 0 else (INK if k == 0 else MUTED))
gap = box(s, 0.5, 5.7, 12.3, 1.1, fill=AMBER, name="Gap")
label_in(gap, "The gap: Gradient Routing needs humans to assign data to regions. Ours lets the model decide.",
         size=22, bold=True, color=NAVY)
notes(s, """[3:00–4:00]  (Presenter 2)
Three lines of existing research touch on this.
Mixture of Experts, used in Switch Transformer and Mixtral, splits the model into expert parts and uses a learned "router" to pick which experts handle each word. But its goal is cheaper computing, not understanding, and the Mixtral authors report that experts do not split up by subject.
Gradient Routing, by Cloud and colleagues in 2024, does localise skills in chosen parts of the model so they can be removed later. But humans have to decide which data goes to which part.
Elastic Weight Consolidation, from 2017, protects the important weights against forgetting, but it needs to know when the task changes.
So here is the gap: nobody lets the model decide the regions itself, without human labels. That is what we try.""")

# =============================================================================== 5. approach
s = new_slide()
title(s, "Our idea: the most active blocks learn", 5)
text(s, 0.6, 1.3, 12.1, 0.8,
     [[("Inside every layer, the ", {}), ("MLP", {"bold": True}),
       (" (the part that stores knowledge) is cut into ", {}), ("8 blocks", {"bold": True}), (".", {})]], size=22)
active = {2, 5}
for i in range(8):
    x = 0.95 + i * 1.5
    b = box(s, x, 2.2, 1.3, 1.1, fill=NAVY if i in active else GREY, radius=0.12)
    label_in(b, f"block {i + 1}", size=16, bold=True, color=WHITE if i in active else MUTED)
    text(s, x - 0.1, 3.35, 1.5, 0.4, "learns" if i in active else "blocked", size=16, bold=i in active,
         color=NAVY if i in active else MUTED, align=PP_ALIGN.CENTER)
text(s, 0.6, 3.85, 12.1, 0.4, "Each batch: only the 2 most active blocks get the learning signal (gradient). Which 2 varies.",
     size=18, color=INK)
steps = [("Measure", "How active is each block?", STORY, LIGHT), ("Select", "Top 2 learn, 6 are blocked", NAVY, LIGHT),
         ("Balance", "A nudge keeps idle blocks alive", AMBER, "FDF3DA")]
for i, (h, d, col, bg) in enumerate(steps):
    x = 0.6 + i * 4.15
    box(s, x, 4.55, 3.9, 1.35, fill=bg)
    n = box(s, x + 0.15, 4.7, 0.5, 0.5, fill=col, shape=MSO_SHAPE.OVAL)
    label_in(n, str(i + 1), size=16, bold=True, color=NAVY if col == AMBER else WHITE)
    text(s, x + 0.8, 4.68, 3.0, 0.55, h, size=22, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.15, 5.3, 3.6, 0.55, d, size=18, color=INK)
q = box(s, 0.6, 6.15, 12.1, 0.7, fill=NAVY)
label_in(q, "Like brain regions: they specialise through use and competition.", size=22, bold=True, color=WHITE)
notes(s, """[4:00–5:00]
Here is our idea. Inside every layer there is a part called the MLP, which stores most of the model's knowledge. We cut it into 8 equal blocks.
For each batch of text, we measure how active each block is. Only the two most active blocks in each layer are allowed to learn from that batch: the learning signal, called the gradient, is blocked for the other six. All eight blocks still do their work when the model reads text. Only learning is restricted.
There is one problem. A block that loses could never come back, so we add a balancing nudge that keeps idle blocks alive.
The inspiration is the brain: regions specialise because they are used, and they compete. Our hope is that a block that wins on math keeps winning on math.
▶ HANDOVER: Presenter 2 hands over to Presenter 3 for slide 6.""")

# =============================================================================== 6. how we test it
s = new_slide()
title(s, "How we test it: switch one block off", 6)
text(s, 0.6, 1.3, 12.1, 0.8, "Like studying brain injury: damage one region and see which ability is lost.", size=22)
tag = box(s, 0.6, 2.2, 6.6, 0.4, fill=AMBER, radius=0.5)
label_in(tag, "ILLUSTRATIVE SKETCH, NOT REAL DATA", size=14, bold=True, color=NAVY)


def sketch(x0, y0, pattern, caption):
    for j, (lab, col) in enumerate([("Story", STORY), ("Math", MATH), ("Code", CODE)]):
        c = box(s, x0 + j * 0.9, y0, 0.8, 0.3, fill=col, radius=0.2)
        label_in(c, lab, size=12, bold=True, color=WHITE)
    for i, row in enumerate(pattern):
        for j, v in enumerate(row):
            box(s, x0 + j * 0.9, y0 + 0.4 + i * 0.4, 0.8, 0.34, fill=v, radius=0.1)
    text(s, x0 - 0.1, y0 + 2.85, 3.0, 0.9, caption, size=16, bold=True, color=NAVY, align=PP_ALIGN.CENTER)


PALE, HOT, FAINT = "F1CFCB", RED, "F8EEEC"
sketch(0.9, 2.8, [[PALE] * 3 for _ in range(6)], "Dense model: every block matters a little for everything")
ideal = [[HOT, FAINT, FAINT], [FAINT, HOT, FAINT], [FAINT, FAINT, HOT], [HOT, FAINT, FAINT], [FAINT, HOT, FAINT],
         [FAINT, FAINT, HOT]]
sketch(4.3, 2.8, ideal, "Specialised model: each block matters for one skill")
# right: reading guide + S scale
box(s, 7.7, 2.2, 5.1, 4.6, fill=LIGHT)
text(s, 7.95, 2.3, 4.7, 1.9,
     ["Each row: one block switched off.", "Colour: how much the error rises for each skill."], size=18)
text(s, 7.95, 3.85, 4.7, 0.5, "Specialization score S", size=22, bold=True, color=NAVY)
for k in range(5):
    shade = ["DCEBFA", "B3D1F3", "7FB0EA", "4D91E0", STORY][k]
    box(s, 7.95 + k * 0.94, 4.5, 0.9, 0.5, fill=shade, radius=0.1)
text(s, 7.9, 5.1, 2.4, 0.9, ["S = 0", "skill shared by all"], size=16, bold=False, color=INK)
text(s, 10.4, 5.1, 2.35, 0.9, ["S = 1", "one skill only"], size=16, bold=False, color=INK, align=PP_ALIGN.RIGHT)
text(s, 7.95, 5.95, 4.7, 0.8, "High S: switching off that block removes one skill.", size=18, bold=True, color=NAVY)
notes(s, """[5:00–6:00]  (Presenter 3)
How do we know whether the blocks specialised? We copy what neurologists do: switch one part off and see what breaks.
We zero out one block at a time and measure how much the model's error rises on each of the three skills: stories, math and code. In the heatmap that we will show shortly, each row is one block, and the colour shows how much worse the model gets.
On the left is an illustrative sketch, not real data. In a normal, dense model, every block matters a little for everything, so the rows look the same. On the right is what a specialised model would look like: each block hurts mainly one skill.
We summarise this in one number, the specialization score S. Zero means a block's skill is shared across all domains. One means the block serves a single domain.""")

# =============================================================================== 7. implementation
s = new_slide()
title(s, "What we built", 7)
stats = [("10.6M", "learned numbers\n(29.9M with word table)"), ("6 × 8 = 48", "layers × blocks to study"),
         ("19 min", "to train on a free\nKaggle T4 GPU")]
for i, (big, small) in enumerate(stats):
    y = 1.45 + i * 1.5
    box(s, 0.5, y, 3.6, 1.35, fill=LIGHT)
    text(s, 0.65, y + 0.05, 3.3, 0.65, big, size=34, bold=True, color=NAVY, font=HEAD)
    text(s, 0.65, y + 0.7, 3.3, 0.65, small.split("\n"), size=14, color=MUTED)
data = [("Stories", "TinyStories", "20.0M tokens", "“Once upon a time there lived a little girl…”", STORY),
        ("Math", "Our own generator", "10.0M tokens", "Q: 347 + 589 = A: 936", MATH),
        ("Code", "CodeParrot (Python)", "10.6M tokens", "def add(a, b): …", CODE)]
for i, (dom, src, tok, ex, col) in enumerate(data):
    x = 4.4 + i * 2.8
    box(s, x, 1.45, 2.65, 4.35, fill=col)
    text(s, x + 0.15, 1.55, 2.35, 0.6, dom, size=26, bold=True, color=WHITE, font=HEAD)
    text(s, x + 0.15, 2.25, 2.35, 0.8, src, size=18, bold=True, color=WHITE)
    text(s, x + 0.15, 3.05, 2.35, 0.5, tok, size=18, color=WHITE)
    ex_box = box(s, x + 0.15, 3.85, 2.35, 1.75, fill=WHITE)
    label_in(ex_box, ex, size=16, color=INK)
text(s, 4.4, 5.85, 8.4, 0.4, "Token = a word piece. 40.6M tokens in total.", size=14, color=MUTED)
for i, t in enumerate(["Python", "PyTorch", "nanoGPT-style GPT", "Kaggle T4 GPU", "GitHub"]):
    w = [1.5, 1.6, 2.6, 2.1, 1.4][i]
    x = 0.5 + sum([1.5, 1.6, 2.6, 2.1, 1.4][:i]) + 0.15 * i
    chip(s, x, 6.35, w, 0.5, t, NAVY, size=16)
t = box(s, 10.35, 6.3, 2.45, 0.6, fill=CODE)
label_in(t, "14 / 14 tests pass", size=18, bold=True, color=WHITE)
notes(s, """[6:00–7:00]
What we built. The model is a small GPT-style language model, written in PyTorch following the nanoGPT design: 6 layers, 384 numbers wide, about 10.6 million learned numbers in the core, or 29.9 million counting the word table. With 8 blocks in each of the 6 layers, there are 48 blocks to study.
It trains on 40.6 million tokens, where a token is a word piece, from three sources: TinyStories for stories, 20 million tokens; our own generator for arithmetic lines like this one, 10 million; and CodeParrot for Python code, 10.6 million.
We trained on a free Kaggle T4 GPU; the baseline took about 19 minutes. Everything is on GitHub with an automated pipeline from data to plots, and all 14 automated tests pass, including tests that prove only the chosen two blocks receive the learning signal.""")

# =============================================================================== 8. results
s = new_slide()
title(s, "Results so far", 8)
picture(s, "results/baseline/loss_curves.png", 0.5, 1.3, w=4.5)
text(s, 0.5, 3.98, 4.5, 0.3, "Loss curves (dense): lower is better", size=12, color=MUTED)
rows8 = [("", "Dense", "Routed"), ("Story loss", "2.120", "2.270"), ("Math loss", "1.467", "1.510"),
         ("Code loss", "2.497", "2.659"), ("Mean S", "0.051", "0.110")]
t8 = s.shapes.add_table(5, 3, Inches(0.5), Inches(4.35), Inches(4.5), Inches(2.0)).table
for j, w in enumerate([1.7, 1.4, 1.4]):
    t8.columns[j].width = Inches(w)
for i, row in enumerate(rows8):
    t8.rows[i].height = Inches(0.4)
    for j, val in enumerate(row):
        c = t8.cell(i, j)
        c.fill.solid()
        c.fill.fore_color.rgb = rgb(NAVY if i == 0 else ("DCEBFA" if i == 4 else (WHITE if i % 2 else LIGHT)))
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        c.margin_top = c.margin_bottom = Inches(0.02)
        p = c.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
        r = p.add_run()
        r.text = val
        r.font.name, r.font.size = BODY, Pt(16)
        r.font.bold = i in (0, 4)
        r.font.color.rgb = rgb(WHITE if i == 0 else INK)
for k, (lab, rel) in enumerate([("Dense", "results/baseline/ablation_heatmap.png"),
                                ("Routed", "results/routed/ablation_heatmap.png")]):
    x = 5.3 + k * 2.2
    text(s, x, 1.25, 2.0, 0.35, lab, size=18, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    picture(s, rel, x, 1.6, h=4.4)
text(s, 5.0, 6.02, 4.7, 0.4, "Rows: blocks switched off. Darker = more damage. Scales differ.", size=12, color=MUTED, align=PP_ALIGN.CENTER)
sm = [("Story", STORY, "“Once upon a time there lived a tidy little girl. One day she went to the park with her mom.”"),
      ("Math", MATH, "Q: 26 - 5 = A: 21"),
      ("Code", CODE, "def add(a, b): o = b.read() … return o\n(looks like code, not valid yet)")]
text(s, 9.75, 1.25, 3.0, 0.35, "Model output (dense)", size=18, bold=True, color=NAVY)
for i, (d, col, ex) in enumerate(sm):
    y = 1.65 + i * 1.5
    box(s, 9.75, y, 3.05, 1.4, fill=LIGHT)
    chip(s, 9.85, y + 0.08, 0.9, 0.28, d, col, size=12)
    text(s, 9.85, y + 0.4, 2.85, 1.0, ex, size=12, color=INK)
tk = box(s, 0.5, 6.45, 12.3, 0.55, fill=NAVY)
label_in(tk, "Takeaway: routing doubled specialization (S 0.05 → 0.11) for a 3–7% loss cost. One run each, so not yet proven.",
         size=17, bold=True, color=WHITE)
notes(s, """[7:00–8:00]
Results so far. Both models trained for 5,000 steps, about 20 minutes each. Top left, the loss curves: lower is better, and the model learns all three domains.
In the table, "dense" is the normal model and "routed" uses our method. Routing costs a little accuracy: story loss goes from 2.12 to 2.27, math from 1.47 to 1.51, code from 2.50 to 2.66, about 3 to 7 percent worse.
The key number is the specialization score: 0.051 for dense, 0.110 for routed. Routing roughly doubled it, and the routed heatmap shows a few dark, domain-specific blocks. The colour scales differ, so compare the shapes, not the shades.
On the right are real samples from the dense model: stories are readable, the math format is learned but answers are often wrong, and the code is not valid yet.
Honest caveat: one run per model, so we cannot yet prove the gap is not luck.
▶ HANDOVER: Presenter 3 hands over to Presenter 4 for slide 9.""")

# =============================================================================== 9. finding + limitations
s = new_slide()
title(s, "A key finding, and what is still open", 9)


def grid(x0, y0, n_dead, caption, heading):
    text(s, x0, y0 - 0.55, 5.9, 0.5, heading, size=22, bold=True, color=NAVY)
    for i in range(16):
        r, c = divmod(i, 8)
        dead = i < n_dead if n_dead else False
        # spread the dead blocks over the grid so it does not look like a sorted list
        order = [0, 3, 5, 8, 10, 11, 13, 14, 15]
        dead = i in order[:n_dead]
        box(s, x0 + c * 0.7, y0 + r * 0.7, 0.6, 0.6, fill=GREY if dead else NAVY, radius=0.15)
    text(s, x0, y0 + 1.5, 5.9, 0.7, caption, size=16, color=INK)


grid(0.6, 2.05, 9, "9 of 16 blocks never selected (small test): the balancing signal was blocked too.",
     "Before: naive masking")
arrow(s, 6.35, 2.7, 7.05, 2.7, color=MUTED, w=3)
grid(7.25, 2.05, 0, "No block left behind. Full run: all 48 blocks used.", "After: balancing signal not masked")
lims = [("Dominant block", "Layer 0, block 6 took 70% of picks; most took 15–31%."),
        ("Mixed batches", "A batch holds stories, math and code, so each pick averages them."),
        ("Optimizer drift", "AdamW still nudges blocks we blocked.")]
text(s, 0.6, 4.3, 12, 0.45, "Known limitations", size=22, bold=True, color=NAVY)
for i, (h, d) in enumerate(lims):
    x = 0.5 + i * 4.2
    box(s, x, 4.85, 3.9, 2.0, fill=LIGHT)
    text(s, x + 0.2, 4.95, 3.5, 0.5, h, size=20, bold=True, color=NAVY)
    text(s, x + 0.2, 5.5, 3.5, 1.3, d, size=18, color=INK)
notes(s, """[8:00–9:00]  (Presenter 4)
An important finding along the way. In our first version, the balancing nudge was blocked too, along with everything else. In a small test, 9 of 16 blocks were never chosen at all. Once a block loses, it gets no learning signal, so it can never win again. We let the balancing signal bypass the block, so idle blocks keep getting nudged, and every block stays in use. In the full run, all 48 blocks were used. We kept the failure as an automated test.
What is still open. First, one block, block 6 in layer 0, was chosen 70 percent of the time, while most blocks were chosen 15 to 31 percent. Second, each batch mixes stories, math and code, so a block's pick averages over the domains. Third, the optimizer we use, AdamW, still slightly moves blocks whose learning signal we blocked, through momentum and weight decay.""")

# =============================================================================== 10. progress + next
s = new_slide()
title(s, "Progress and next steps", 10)
mods = [("Data pipeline", "Completed", NAVY, "41M tokens, 3 domains"),
        ("Baseline training", "Completed", NAVY, "5,000 steps on a T4"),
        ("Routing module", "Completed", NAVY, "14 tests, 1 full run"),
        ("Evaluation harness", "Ongoing", AMBER, "Ablation done, probes pending")]
for i, (m, st, col, d) in enumerate(mods):
    y = 1.45 + i * 1.05
    box(s, 0.5, y, 7.0, 0.95, fill=LIGHT)
    text(s, 0.7, y + 0.04, 3.4, 0.5, m, size=20, bold=True, color=NAVY)
    text(s, 0.7, y + 0.5, 3.6, 0.4, d, size=14, color=MUTED)
    chip(s, 5.6, y + 0.22, 1.7, 0.5, st, col, color=WHITE if col == NAVY else NAVY, size=16)
text(s, 7.9, 1.4, 4.9, 0.5, "Next steps", size=22, bold=True, color=NAVY)
steps10 = ["Repeat with several seeds", "Human-assigned regions baseline", "Fix dominant block and drift",
           "Consolidation phase", "Domain probes", "Unlearning and continual learning"]
for i, t in enumerate(steps10):
    y = 1.95 + i * 0.62
    c = box(s, 7.9, y, 0.45, 0.45, fill=STORY, shape=MSO_SHAPE.OVAL)
    label_in(c, str(i + 1), size=16, bold=True, color=WHITE)
    text(s, 8.5, y - 0.03, 4.4, 0.55, t, size=18, color=INK, anchor=MSO_ANCHOR.MIDDLE)
ty = box(s, 0.5, 5.85, 12.3, 1.0, fill=NAVY)
label_in(ty, "Thank you  ·  Questions?", size=34, bold=True, color=WHITE, font=HEAD)
notes(s, """[9:00–10:00]
To summarise our progress. The data pipeline and baseline training are complete. The routing module is complete too: implemented, tested, and run once at full scale. The evaluation harness is ongoing: the switch-one-block-off test is done, but the probes that read what each block encodes are still pending.
Next, we will repeat the runs with several seeds, to see whether the doubling of S is real. Then a baseline where humans assign the regions, which is our upper bound. We will also fix the dominant block and the optimizer drift, add a consolidation phase, build the domain probes, and run unlearning and continual-learning experiments.
Thank you. We are happy to take your questions.""")

prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides))
